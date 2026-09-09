#!/usr/bin/env python3
"""Сборка PPTX-копии деки из talk/deck.html.

Копия нужна там, где организаторы принимают только PowerPoint. Слайды собираются
нативными объектами — текст остаётся текстом, таблицы таблицами, — поэтому копию
можно править в PowerPoint. Пиксельного совпадения с HTML-декой нет и не требуется:
источник правды — deck.html, pptx пересобирается из него.

    python3 talk/deck2pptx.py            # → talk/deck.pptx

Скриншоты из talk/shots/ подставляются, если файлы есть; иначе на их месте
остаётся рамка с описанием кадра — та же, что в HTML-деке.
"""
import os
import re
import sys
from html import unescape
from html.parser import HTMLParser

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "deck.html")
SHOTS = os.path.join(HERE, "shots")
OUT = os.path.join(HERE, "deck.pptx")

# Палитра деки (:root в deck.html) — держать синхронной с ней.
INK = RGBColor(0x14, 0x17, 0x1F)
PANEL = RGBColor(0x1D, 0x22, 0x2D)
PANEL2 = RGBColor(0x23, 0x29, 0x37)
LINE = RGBColor(0x2E, 0x35, 0x42)
TEXT = RGBColor(0xE9, 0xEC, 0xF2)
MUTED = RGBColor(0x98, 0xA1, 0xB3)
DIM = RGBColor(0x5C, 0x65, 0x77)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
AMBER_BG = RGBColor(0x2A, 0x25, 0x1C)
OK = RGBColor(0x4F, 0xC3, 0x8A)
BAD = RGBColor(0xE0, 0x56, 0x4F)

SANS = "Avenir Next"
MONO = "Menlo"

W, H = Inches(13.333), Inches(7.5)
PAD = Inches(0.62)
BODY_TOP = Inches(1.62)
BODY_BOTTOM = Inches(7.0)
GAP = Inches(0.26)

CLASS_COLOR = {"hl": AMBER, "acc": AMBER, "good": OK, "vmark": OK,
               "badc": BAD, "xmark": BAD, "dim": DIM, "muted": MUTED}


def E(value):
    """EMU обязан быть целым: дробная координата даёт невалидный OOXML.

    Деление высоты колонки между атомами почти всегда даёт float, поэтому
    все размеры прогоняются через это до попадания в фигуру.
    """
    return Emu(int(round(float(value))))


# --- разбор HTML -----------------------------------------------------------


class Node:
    __slots__ = ("tag", "attrs", "kids", "parent")

    def __init__(self, tag, attrs=None, parent=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.kids = []
        self.parent = parent

    @property
    def cls(self):
        return self.attrs.get("class", "").split()

    def has(self, name):
        return name in self.cls

    def find_all(self, pred):
        found = []
        for kid in self.kids:
            if isinstance(kid, Node):
                if pred(kid):
                    found.append(kid)
                found.extend(kid.find_all(pred))
        return found

    def text(self):
        out = []
        for kid in self.kids:
            out.append(kid if isinstance(kid, str) else kid.text())
        return re.sub(r"\s+", " ", "".join(out)).strip()


VOID = {"br", "img", "meta", "link", "hr", "input"}


class TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag, dict(attrs), self.stack[-1])
        self.stack[-1].kids.append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        if data.strip():
            self.stack[-1].kids.append(data)


def parse_slides():
    html = open(DECK, encoding="utf-8").read()
    html = html[html.index('<template id="slides">'):]
    builder = TreeBuilder()
    builder.feed(html)
    return builder.root.find_all(lambda n: n.tag == "section" and n.has("slide"))


# --- сбор inline-текста в runs ---------------------------------------------


def runs_of(node, bold=False, color=None, mono=False):
    """Плоский список (текст, жирный, цвет, моноширинный) с учётом вложенных тегов."""
    out = []
    for kid in node.kids:
        if isinstance(kid, str):
            out.append((unescape(kid), bold, color, mono))
            continue
        if kid.tag == "br":
            out.append(("\n", bold, color, mono))
            continue
        if kid.tag in ("aside", "template"):
            continue
        kid_color = color
        for name in kid.cls:
            if name in CLASS_COLOR:
                kid_color = CLASS_COLOR[name]
        out.extend(runs_of(kid,
                           bold=bold or kid.tag in ("b", "strong"),
                           color=kid_color,
                           mono=mono or kid.tag in ("code", "pre")))
    return out


def collapse(runs):
    """Схлопывает пробелы, сохраняя явные переводы строки."""
    out = []
    for text, bold, color, mono in runs:
        text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        if not text:
            continue
        if out and out[-1][1:] == (bold, color, mono):
            out[-1] = (out[-1][0] + text, bold, color, mono)
        else:
            out.append((text, bold, color, mono))
    if out:
        out[0] = (out[0][0].lstrip(), *out[0][1:])
        out[-1] = (out[-1][0].rstrip(), *out[-1][1:])
    return [r for r in out if r[0]]


# --- примитивы отрисовки ---------------------------------------------------


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(E(left), E(top), E(width), E(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    return frame


def write(frame, runs, size, color=TEXT, font=SANS, bold=False,
          align=PP_ALIGN.LEFT, spacing=1.25, first=True):
    """Пишет runs в текстовый фрейм; \n внутри run начинает новый абзац."""
    para = frame.paragraphs[0] if first else frame.add_paragraph()
    para.alignment = align
    para.line_spacing = spacing
    for text, is_bold, run_color, mono in runs:
        chunks = text.split("\n")
        for i, chunk in enumerate(chunks):
            if i:
                para = frame.add_paragraph()
                para.alignment = align
                para.line_spacing = spacing
            if not chunk:
                continue
            run = para.add_run()
            run.text = chunk
            run.font.size = Pt(size)
            run.font.name = MONO if mono else font
            run.font.bold = bold or is_bold
            run.font.color.rgb = run_color or color
    return para


def panel(slide, left, top, width, height, fill=PANEL, border=LINE, dashed=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(left), E(top), E(width), E(height))
    shape.adjustments[0] = 0.04
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(1)
    if dashed:
        shape.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False
    shape.text_frame.word_wrap = True
    return shape


def label(slide, left, top, width, text, color=DIM):
    frame = textbox(slide, left, top, width, Inches(0.22))
    write(frame, [(text.upper(), False, None, False)], 9, color=color, font=MONO, spacing=1.0)
    return Inches(0.28)


# --- атомы контента --------------------------------------------------------


def weight(node):
    """Грубая оценка «сколько места просит атом» — для дележа высоты колонки."""
    if node.has("shot"):
        return 3.0
    if node.tag == "table":
        return 1.0 + 0.55 * len(node.find_all(lambda n: n.tag == "tr"))
    if node.has("chain"):
        return 1.2 * len(node.find_all(lambda n: n.has("node")))
    if node.has("rate"):
        return 1.4
    if node.has("thesis"):
        return 2.0
    return max(1.2, len(node.text()) / 150.0)


def draw_shot(slide, node, left, top, width, height):
    tag = next((k.text() for k in node.kids if isinstance(k, Node) and k.has("shot-tag")), "")
    desc = next((k.text() for k in node.kids if isinstance(k, Node) and k.has("shot-desc")), "")
    src = node.attrs.get("data-src", "")
    path = os.path.join(HERE, src) if src else ""
    if path and os.path.exists(path):
        slide.shapes.add_picture(path, E(left), E(top), width=E(width))
        return
    panel(slide, left, top, width, height, fill=AMBER_BG, border=AMBER, dashed=True)
    frame = textbox(slide, left + Inches(0.16), top + Inches(0.14),
                    width - Inches(0.32), height - Inches(0.28))
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    write(frame, [(tag.upper(), False, None, False)], 9, color=AMBER, font=MONO,
          align=PP_ALIGN.CENTER, spacing=1.0)
    write(frame, [(desc, False, None, False)], 9.5, color=MUTED,
          align=PP_ALIGN.CENTER, spacing=1.2, first=False)


def draw_table(slide, node, left, top, width, height):
    caption = node.find_all(lambda n: n.tag == "caption")
    if caption:
        top += label(slide, left, top, width, caption[0].text(), color=AMBER)
        height -= Inches(0.28)
    rows = node.find_all(lambda n: n.tag == "tr")
    cols = max(len(r.find_all(lambda n: n.tag in ("td", "th"))) for r in rows)
    shape = slide.shapes.add_table(len(rows), cols, E(left), E(top), E(width), E(height))
    table = shape.table
    first_wide = node.has("ev")
    if first_wide and cols > 1:
        table.columns[0].width = E(width * 0.42)
        for i in range(1, cols):
            table.columns[i].width = E(width * 0.58 / (cols - 1))
    for r, row in enumerate(rows):
        cells = row.find_all(lambda n: n.tag in ("td", "th"))
        is_head = any(c.tag == "th" for c in cells)
        is_calls = row.has("calls")
        table.rows[r].height = Inches(0.26 if is_calls else 0.32)
        for c in range(cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if c >= len(cells):
                cell.text_frame.text = ""
                continue
            src = cells[c]
            color = DIM if (is_head or is_calls) else TEXT
            for name in src.cls:
                if name == "g":
                    color = OK
                elif name == "b":
                    color = BAD
            frame = cell.text_frame
            frame.word_wrap = True
            write(frame, collapse(runs_of(src)) or [("", False, None, False)],
                  9 if (is_head or is_calls) else 10.5,
                  color=color, font=MONO if first_wide else SANS,
                  bold=is_head, spacing=1.0,
                  align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)


def draw_chain(slide, node, left, top, width, height):
    nodes = [k for k in node.kids if isinstance(k, Node) and k.has("node")]
    arrows = [k for k in node.kids if isinstance(k, Node) and k.has("arr")]
    n = len(nodes)
    arrow_h = Inches(0.3)
    box_h = (height - arrow_h * len(arrows)) / max(n, 1)
    y = top
    for i, item in enumerate(nodes):
        em = item.find_all(lambda k: k.tag == "em")
        title = "".join(k for k in item.kids if isinstance(k, str)).strip()
        border = AMBER if "rgba(232,163,61" in item.attrs.get("style", "") else LINE
        panel(slide, left, y, width, box_h, fill=PANEL2, border=border)
        frame = textbox(slide, left + Inches(0.12), y + Inches(0.08),
                        width - Inches(0.24), box_h - Inches(0.16))
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        write(frame, [(title, True, None, False)], 11, font=MONO,
              align=PP_ALIGN.CENTER, spacing=1.0)
        if em:
            write(frame, [(em[0].text(), False, None, False)], 8.5, color=MUTED,
                  align=PP_ALIGN.CENTER, spacing=1.05, first=False)
        y += box_h
        if i < len(arrows):
            frame = textbox(slide, left, y, width, arrow_h)
            write(frame, [(arrows[i].text(), False, None, False)], 9, color=AMBER,
                  font=MONO, align=PP_ALIGN.CENTER, spacing=1.0)
            y += arrow_h


def draw_rate(slide, node, left, top, width, height):
    bars = [k for k in node.kids if isinstance(k, Node) and k.tag == "i"]
    gap = Inches(0.04)
    bar_w = (width - gap * (len(bars) - 1)) / len(bars)
    for i, bar in enumerate(bars):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, E(left + i * (bar_w + gap)), E(top), E(bar_w), Inches(0.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = BAD if bar.has("f") else OK
        shape.line.fill.background()
        shape.shadow.inherit = False


def draw_pre(slide, node, left, top, width, height):
    panel(slide, left, top, width, height)
    frame = textbox(slide, left + Inches(0.14), top + Inches(0.11),
                    width - Inches(0.28), height - Inches(0.22))
    write(frame, collapse(runs_of(node)), 10, font=MONO, spacing=1.2)


def draw_card(slide, node, left, top, width, height):
    panel(slide, left, top, width, height)
    inner_l, inner_w = left + Inches(0.18), width - Inches(0.36)
    y = top + Inches(0.14)
    for kid in node.kids:
        if not isinstance(kid, Node):
            continue
        if kid.has("lbl"):
            y += label(slide, inner_l, y, inner_w, kid.text())
        elif kid.tag == "div" and "height:1px" in kid.attrs.get("style", ""):
            rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(inner_l), E(y), E(inner_w), Pt(1))
            rule.fill.solid()
            rule.fill.fore_color.rgb = LINE
            rule.line.fill.background()
            rule.shadow.inherit = False
            y += Inches(0.14)
        else:
            runs = collapse(runs_of(kid))
            if not runs:
                continue
            size = 9.5 if kid.has("small") else 12
            color = MUTED if kid.has("small") else TEXT
            lines = sum(t.count("\n") for t, *_ in runs) + 1
            chars = sum(len(t) for t, *_ in runs)
            box_h = Inches(0.02) + Pt(size * 1.5) * max(lines, chars / 52.0)
            frame = textbox(slide, inner_l, y, inner_w, box_h)
            write(frame, runs, size, color=color, spacing=1.3)
            y += box_h + Inches(0.08)


def draw_list(slide, node, left, top, width, height, columns=1):
    items = [k for k in node.kids if isinstance(k, Node) and k.tag == "li"]
    size = 11 if node.has("check") else 12.5
    per = (len(items) + columns - 1) // columns
    col_w = (width - GAP * (columns - 1)) / columns
    for c in range(columns):
        chunk = items[c * per:(c + 1) * per]
        if not chunk:
            continue
        frame = textbox(slide, left + c * (col_w + GAP), top, col_w, height)
        for i, item in enumerate(chunk):
            runs = collapse(runs_of(item))
            write(frame, [("•  ", False, AMBER, False)] + runs, size,
                  spacing=1.3, first=(i == 0))
            frame.paragraphs[-1].space_after = Pt(7)


def draw_banner(slide, node, left, top, width):
    height = Inches(0.52)
    shape = panel(slide, left, top, width, height, fill=AMBER_BG, border=AMBER)
    frame = shape.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Inches(0.14)
    write(frame, collapse(runs_of(node)), 12.5, bold=True, align=PP_ALIGN.CENTER, spacing=1.1)
    return height


def draw_fails(slide, node, left, top, width, height):
    cards = [k for k in node.kids if isinstance(k, Node) and k.has("fail")]
    col_w = (width - GAP * (len(cards) - 1)) / len(cards)
    for i, card in enumerate(cards):
        x = left + i * (col_w + GAP)
        panel(slide, x, top, col_w, height)
        frame = textbox(slide, x + Inches(0.16), top + Inches(0.14),
                        col_w - Inches(0.32), height - Inches(0.28))
        first = True
        for kid in card.kids:
            if not isinstance(kid, Node):
                continue
            if kid.has("n"):
                write(frame, [(kid.text().upper(), False, None, False)], 9,
                      color=AMBER, font=MONO, spacing=1.0, first=first)
            elif kid.tag == "b":
                write(frame, [(kid.text(), True, None, False)], 13, spacing=1.15, first=first)
            else:
                write(frame, collapse(runs_of(kid)), 10, color=MUTED, spacing=1.25, first=first)
            first = False


def draw_atom(slide, node, left, top, width, height):
    if node.has("shot"):
        draw_shot(slide, node, left, top, width, height)
    elif node.tag == "table":
        draw_table(slide, node, left, top, width, height)
    elif node.has("chain"):
        draw_chain(slide, node, left, top, width, height)
    elif node.has("rate"):
        draw_rate(slide, node, left, top, width, height)
    elif node.tag == "pre":
        draw_pre(slide, node, left, top, width, height)
    elif node.has("card") or node.has("col") and node.has("card"):
        draw_card(slide, node, left, top, width, height)
    elif node.tag == "ul":
        draw_list(slide, node, left, top, width, height,
                  columns=2 if node.has("check") else 1)
    elif node.has("fails"):
        draw_fails(slide, node, left, top, width, height)
    elif node.has("thesis"):
        frame = textbox(slide, left, top, width, height)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        write(frame, collapse(runs_of(node)), 19, bold=True, spacing=1.3)
    elif node.has("lbl"):
        label(slide, left, top, width, node.text())
    else:
        runs = collapse(runs_of(node))
        if runs:
            frame = textbox(slide, left, top, width, height)
            size = 9.5 if node.has("small") else 12.5
            write(frame, runs, size, color=MUTED if node.has("small") else TEXT, spacing=1.3)


def atoms_of(container):
    """Значимые прямые потомки — то, что рисуется вертикальным стеком."""
    out = []
    for kid in container.kids:
        if not isinstance(kid, Node) or kid.tag == "aside":
            continue
        if kid.tag == "div" and not kid.cls and not kid.kids:
            continue
        out.append(kid)
    return out


def stack(slide, container, left, top, width, height):
    items = atoms_of(container)
    if not items:
        return
    weights = [weight(i) for i in items]
    total = sum(weights)
    free = height - GAP * (len(items) - 1)
    y = top
    for item, w in zip(items, weights):
        h = free * (w / total)
        draw_atom(slide, item, left, y, width, h)
        y += h + GAP


# --- сборка слайда ---------------------------------------------------------


def build_slide(prs, section, index, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = INK

    body_w = W - PAD * 2
    eyebrow = section.find_all(lambda n: n.has("eyebrow"))
    title = section.find_all(lambda n: n.tag == "h1")
    banner = section.find_all(lambda n: n.has("banner"))

    if eyebrow:
        spans = [s.text() for s in eyebrow[0].find_all(lambda n: n.tag == "span")]
        frame = textbox(slide, PAD, Inches(0.42), body_w, Inches(0.24))
        write(frame, [((spans[0] if spans else "").upper(), False, None, False)],
              9.5, color=DIM, font=MONO, spacing=1.0)
        if len(spans) > 1:
            frame2 = textbox(slide, PAD, Inches(0.42), body_w, Inches(0.24))
            write(frame2, [(spans[1].upper(), False, None, False)], 9.5, color=DIM,
                  font=MONO, align=PP_ALIGN.RIGHT, spacing=1.0)

    top = BODY_TOP
    if section.find_all(lambda n: n.has("title-wrap")):
        top = Inches(2.3)
        wrap = section.find_all(lambda n: n.has("title-wrap"))[0]
        head = wrap.find_all(lambda n: n.tag == "h1")[0]
        frame = textbox(slide, PAD, top, Inches(9.5), Inches(1.9))
        write(frame, collapse(runs_of(head)), 42, bold=True, spacing=1.08)
        sub = wrap.find_all(lambda n: n.has("sub"))
        if sub:
            frame = textbox(slide, PAD, top + Inches(2.0), Inches(9.0), Inches(0.8))
            write(frame, collapse(runs_of(sub[0])), 20, color=MUTED, spacing=1.3)
        meta = section.find_all(lambda n: n.has("title-meta"))
        if meta:
            spans = meta[0].find_all(lambda n: n.tag == "span")
            frame = textbox(slide, PAD, Inches(6.5), body_w, Inches(0.3))
            write(frame, [(spans[0].text(), False, None, False)], 10, color=DIM,
                  font=MONO, spacing=1.0)
    else:
        if title:
            frame = textbox(slide, PAD, Inches(0.92), body_w, Inches(0.7))
            write(frame, collapse(runs_of(title[0])), 23, bold=True, spacing=1.14)

        bottom = BODY_BOTTOM
        if banner:
            bottom -= Inches(0.68)
        cols = section.find_all(lambda n: n.has("cols"))
        if cols:
            columns = [c for c in cols[0].kids if isinstance(c, Node) and c.has("col")]
            col_w = (body_w - GAP * (len(columns) - 1)) / len(columns)
            for i, col in enumerate(columns):
                x = PAD + i * (col_w + GAP)
                if col.has("card"):
                    draw_card(slide, col, x, top, col_w, bottom - top)
                else:
                    stack(slide, col, x, top, col_w, bottom - top)
        else:
            container = Node("div")
            container.kids = [k for k in section.kids
                              if isinstance(k, Node)
                              and not k.has("eyebrow") and k.tag not in ("h1", "aside")
                              and not k.has("banner")]
            stack(slide, container, PAD, top, body_w, bottom - top)

        if banner:
            draw_banner(slide, banner[0], PAD, bottom + Inches(0.16), body_w)

    frame = textbox(slide, W - PAD - Inches(1.2), Inches(7.05), Inches(1.2), Inches(0.25))
    write(frame, [("%d / %d" % (index, total), False, None, False)], 9,
          color=DIM, font=MONO, align=PP_ALIGN.RIGHT, spacing=1.0)

    notes = section.find_all(lambda n: n.tag == "aside")
    if notes:
        slide.notes_slide.notes_text_frame.text = notes[0].text()
    return slide


def main():
    sections = parse_slides()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for i, section in enumerate(sections, 1):
        build_slide(prs, section, i, len(sections))
    prs.save(OUT)
    have = sum(1 for f in ("shot-1.png", "shot-2.png", "shot-3.png", "shot-4.png",
                           "shot-5.png", "shot-6.png")
               if os.path.exists(os.path.join(SHOTS, f)))
    print("собрано слайдов: %d → %s" % (len(sections), OUT))
    print("скринов подставлено: %d из 6 (остальные — рамки с описанием кадра)" % have)
    return 0


if __name__ == "__main__":
    sys.exit(main())
